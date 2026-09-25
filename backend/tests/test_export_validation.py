"""Iteration 9 – Validasi Periode / Tahun untuk endpoint export.

Endpoint diuji:
- GET /api/budgets/export?period=YYYY-MM
- GET /api/budgets/export-annual?year=YYYY
- GET /api/budgets/export-range?start=&end=
"""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    from pathlib import Path
    for line in Path("/app/frontend/.env").read_text().splitlines():
        if line.startswith("REACT_APP_BACKEND_URL="):
            BASE_URL = line.split("=", 1)[1].strip().rstrip("/")
API = f"{BASE_URL}/api"

SUPER = {"email": "nashoharizal@gmail.com", "password": "SIKPPBJ2026"}
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MSG_PERIOD = "Periode tidak valid. Gunakan format YYYY-MM (contoh: 2026-09)."
MSG_YEAR = "Tahun tidak valid. Masukkan tahun antara 2000 dan 2100."
MSG_START_INVALID = "Periode awal tidak valid"
MSG_START_AFTER = "Periode awal harus lebih awal atau sama dengan periode akhir"


@pytest.fixture(scope="module")
def sess():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json=SUPER, timeout=15)
    assert r.status_code == 200
    return s


# -------------------- MONTHLY EXPORT VALIDATION --------------------

@pytest.mark.parametrize("period", ["2026-13", "bad", "2026-9"])
def test_export_monthly_invalid_period_400(sess, period):
    r = sess.get(f"{API}/budgets/export", params={"period": period}, timeout=15)
    assert r.status_code == 400, f"got {r.status_code} body={r.text[:200]}"
    body = r.json()
    detail = body.get("detail") or body.get("message") or ""
    assert detail == MSG_PERIOD, f"unexpected detail: {detail!r}"


def test_export_monthly_valid_period_200(sess):
    r = sess.get(f"{API}/budgets/export", params={"period": "2026-09"}, timeout=30)
    assert r.status_code == 200
    assert XLSX_MEDIA in r.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(r.content))
    assert "Anggaran vs Realisasi" in wb.sheetnames


# -------------------- ANNUAL EXPORT VALIDATION --------------------

@pytest.mark.parametrize("year", ["1800", "3000"])
def test_export_annual_invalid_year_400(sess, year):
    r = sess.get(f"{API}/budgets/export-annual", params={"year": year}, timeout=15)
    assert r.status_code == 400, f"got {r.status_code} body={r.text[:200]}"
    body = r.json()
    detail = body.get("detail") or body.get("message") or ""
    assert detail == MSG_YEAR, f"unexpected detail: {detail!r}"


def test_export_annual_valid_year_200(sess):
    r = sess.get(f"{API}/budgets/export-annual", params={"year": "2026"}, timeout=30)
    assert r.status_code == 200
    assert XLSX_MEDIA in r.headers.get("content-type", "")


# -------------------- RANGE EXPORT VALIDATION --------------------

def test_export_range_invalid_start_format_400(sess):
    r = sess.get(f"{API}/budgets/export-range",
                 params={"start": "bad", "end": "2026-09"}, timeout=15)
    assert r.status_code == 400
    detail = (r.json().get("detail") or "")
    assert MSG_START_INVALID in detail, f"unexpected detail: {detail!r}"


def test_export_range_start_after_end_400(sess):
    r = sess.get(f"{API}/budgets/export-range",
                 params={"start": "2026-09", "end": "2026-07"}, timeout=15)
    assert r.status_code == 400
    detail = (r.json().get("detail") or "")
    assert MSG_START_AFTER in detail, f"unexpected detail: {detail!r}"


def test_export_range_valid_200(sess):
    r = sess.get(f"{API}/budgets/export-range",
                 params={"start": "2026-07", "end": "2026-09"}, timeout=60)
    assert r.status_code == 200
    assert XLSX_MEDIA in r.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(r.content))
    assert "Ringkasan" in wb.sheetnames
