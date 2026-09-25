"""Tests for Excel export (monthly/annual/range) and upload/download flow.

Review request iteration_7:
- GET /api/budgets/export?period=YYYY-MM  -> xlsx w/ 'Anggaran vs Realisasi'
- GET /api/budgets/export-annual?year=YYYY -> xlsx w/ 'Tahunan {year}'
- GET /api/budgets/export-range?start&end -> xlsx w/ 'Ringkasan' + per-month sheets, 400 on invalid
- POST /api/upload (png/jpeg/pdf) -> {name,storage_path,content_type,url}; >10MB rejected
- GET /api/files/{storage_path} returns file bytes
"""
import io
import os
import base64
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    from pathlib import Path
    for line in Path("/app/frontend/.env").read_text().splitlines():
        if line.startswith("REACT_APP_BACKEND_URL="):
            BASE_URL = line.split("=", 1)[1].strip().rstrip("/")
API = f"{BASE_URL}/api"

SUPER = {"email": "nashoharizal@gmail.com", "password": "SIKPPBJ2026"}
KEU = {"email": "keuangan@sbb.co.id", "password": "keuangan123"}

PERIOD = "2026-09"
YEAR = "2026"

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8/5+hHgAHggJ/PchI7wAAAABJRU5ErkJggg=="
)


@pytest.fixture(scope="module")
def sess():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json=SUPER, timeout=15)
    assert r.status_code == 200, f"login: {r.status_code} {r.text}"
    return s


@pytest.fixture(scope="module")
def keu_sess():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json=KEU, timeout=15)
    assert r.status_code == 200
    return s


# -------------------- EXPORT MONTHLY --------------------

def test_export_monthly_ok(sess):
    r = sess.get(f"{API}/budgets/export", params={"period": PERIOD}, timeout=30)
    assert r.status_code == 200, r.text[:200]
    ct = r.headers.get("content-type", "")
    assert XLSX_MEDIA in ct, f"unexpected content-type: {ct}"
    assert r.content[:2] == b"PK", "not a zip/xlsx file"
    wb = load_workbook(io.BytesIO(r.content), data_only=False)
    assert "Anggaran vs Realisasi" in wb.sheetnames
    ws = wb["Anggaran vs Realisasi"]
    # some cells and at least a couple of data rows
    assert ws.max_row >= 3
    # find total row - period should appear in header area
    all_text = "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )
    # sheet uses Indonesian month name ("September 2026") rather than "YYYY-MM"
    assert "September 2026" in all_text or PERIOD in all_text
    # ensure key headers present
    assert "Pagu" in all_text and "Realisasi" in all_text
    assert "TOTAL" in all_text


# -------------------- EXPORT ANNUAL --------------------

def test_export_annual_ok(sess):
    r = sess.get(f"{API}/budgets/export-annual", params={"year": YEAR}, timeout=30)
    assert r.status_code == 200, r.text[:200]
    assert XLSX_MEDIA in r.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(r.content))
    sheet = f"Tahunan {YEAR}"
    assert sheet in wb.sheetnames, f"sheets={wb.sheetnames}"
    ws = wb[sheet]
    # 12 months rekap → at least a decent number of rows
    assert ws.max_row >= 12


def test_export_annual_bad_year(sess):
    r = sess.get(f"{API}/budgets/export-annual", params={"year": "abcd"}, timeout=15)
    assert r.status_code in (400, 422)


# -------------------- EXPORT RANGE --------------------

def test_export_range_ok(sess):
    r = sess.get(f"{API}/budgets/export-range",
                 params={"start": "2026-08", "end": "2026-10"}, timeout=60)
    assert r.status_code == 200, r.text[:200]
    assert XLSX_MEDIA in r.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(r.content))
    assert "Ringkasan" in wb.sheetnames
    # should have at least 3 monthly sheets (Aug, Sep, Oct)
    monthly = [s for s in wb.sheetnames if s != "Ringkasan"]
    assert len(monthly) >= 3, f"expected >=3 monthly sheets, got {wb.sheetnames}"


def test_export_range_start_after_end_400(sess):
    r = sess.get(f"{API}/budgets/export-range",
                 params={"start": "2026-10", "end": "2026-08"}, timeout=15)
    assert r.status_code == 400


def test_export_range_invalid_format_400(sess):
    r = sess.get(f"{API}/budgets/export-range",
                 params={"start": "2026-99", "end": "2026-08"}, timeout=15)
    assert r.status_code in (400, 422)


# -------------------- UPLOAD / DOWNLOAD --------------------

def test_upload_png_and_fetch(sess):
    files = {"file": ("nota.png", io.BytesIO(PNG_BYTES), "image/png")}
    r = sess.post(f"{API}/upload", files=files, timeout=30)
    assert r.status_code == 200, r.text
    d = r.json()
    for k in ("name", "storage_path", "content_type", "url"):
        assert k in d, f"missing {k}"
    assert d["name"] == "nota.png"
    assert d["content_type"] == "image/png"
    assert d["storage_path"].startswith("sbb-keuangan/uploads/"), d["storage_path"]
    assert d["url"].startswith("/api/files/")

    # fetch back via cookie session
    r2 = sess.get(f"{BASE_URL}{d['url']}", timeout=30)
    assert r2.status_code == 200
    assert r2.headers.get("content-type", "").startswith("image/")
    assert len(r2.content) >= 1


def test_upload_pdf_ok(sess):
    # minimal PDF header
    pdf = b"%PDF-1.4\n%EOF\n"
    files = {"file": ("nota.pdf", io.BytesIO(pdf), "application/pdf")}
    r = sess.post(f"{API}/upload", files=files, timeout=30)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["content_type"] == "application/pdf"


def test_upload_too_large_400(sess):
    big = b"x" * (10 * 1024 * 1024 + 10)  # >10MB
    files = {"file": ("big.bin", io.BytesIO(big), "application/octet-stream")}
    r = sess.post(f"{API}/upload", files=files, timeout=60)
    assert r.status_code == 400, f"expected 400 got {r.status_code} {r.text[:200]}"


# -------------------- ATTACHMENT PERSISTED ON DOC --------------------

def test_document_with_attachment_persists(sess):
    # upload
    up = sess.post(
        f"{API}/upload",
        files={"file": ("bukti.png", io.BytesIO(PNG_BYTES), "image/png")},
        timeout=30,
    ).json()

    body = {
        "doc_type": "PPBJ",
        "unit_kerja": "TEST_ATTACH",
        "kegiatan": "TEST_attach",
        "keterangan": "TEST_attach",
        "tanggal": "2026-09-20",
        "items": [{"uraian": "X", "kuantitas": 1, "satuan": "unit",
                   "harga_estimasi": 1000, "total": 1000}],
        "attachments": [{
            "name": up["name"],
            "storage_path": up["storage_path"],
            "content_type": up["content_type"],
            "url": up["url"],
        }],
    }
    r = sess.post(f"{API}/documents", json=body, timeout=15)
    assert r.status_code == 200, r.text
    did = r.json()["id"]
    g = sess.get(f"{API}/documents/{did}", timeout=15).json()
    atts = g.get("attachments") or []
    assert len(atts) == 1
    assert atts[0]["storage_path"] == up["storage_path"]
    assert atts[0]["url"].startswith("/api/files/")
    # cleanup
    sess.delete(f"{API}/documents/{did}")
